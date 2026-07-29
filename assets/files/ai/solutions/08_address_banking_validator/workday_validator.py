"""
workday_validator.py
Core validation engine for Workday address and banking data.
"""

import re
import os
import tempfile
from typing import Optional

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None


# ---------------------------------------------------------------------------
# ADDRESS RULES
# ---------------------------------------------------------------------------

ADDRESS_RULES = {
    "United States of America": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Optional", "postal_code": "Required"
    },
    "United Kingdom": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Australia": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Canada": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Optional", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Germany": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "France": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "India": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Japan": {
        "address_line_1": "Required", "address_line_2": "Not Allowed", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Optional", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Singapore": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Brazil": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Optional", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Netherlands": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Italy": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Spain": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Optional", "postal_code": "Required"
    },
    "Switzerland": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Sweden": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Belgium": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Mexico": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Required", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Argentina": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Optional", "postal_code": "Required"
    },
    "China": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Optional", "city_subdivision_1": "Optional", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Korea, Republic of": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Optional", "postal_code": "Optional"
    },
    "Hong Kong": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Not Allowed"
    },
    "Denmark": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Optional", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Optional", "postal_code": "Required"
    },
    "Norway": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Not Allowed",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Finland": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Not Allowed", "city_subdivision_1": "Required", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Austria": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Portugal": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Optional", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Ireland": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Optional", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Optional"
    },
    "Luxembourg": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Taiwan": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Thailand": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Not Allowed", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Required", "postal_code": "Required"
    },
    "New Zealand": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Optional", "postal_code": "Required"
    },
    "Israel": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "South Africa": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Optional", "city_subdivision_1": "Required", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Colombia": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Optional"
    },
    "Morocco": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
        "city": "Optional", "city_subdivision_1": "Not Allowed", "region": "Optional",
        "region_subdivision_1": "Not Allowed", "postal_code": "Required"
    },
    "Cayman Islands": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Required", "city_subdivision_1": "Not Allowed", "region": "Not Allowed",
        "region_subdivision_1": "Not Allowed", "postal_code": "Optional"
    },
    "United Arab Emirates": {
        "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Not Allowed",
        "city": "Not Allowed", "city_subdivision_1": "Not Allowed", "region": "Required",
        "region_subdivision_1": "Not Allowed", "postal_code": "Not Allowed"
    },
}

_ADDRESS_RULES_GENERIC = {
    "address_line_1": "Required", "address_line_2": "Optional", "address_line_3": "Optional",
    "city": "Required", "city_subdivision_1": "Optional", "region": "Optional",
    "region_subdivision_1": "Optional", "postal_code": "Optional"
}



# ---------------------------------------------------------------------------
# POSTAL CODE PATTERNS
# ---------------------------------------------------------------------------

POSTAL_CODE_PATTERNS = {
    "United States of America": (r"^\d{5}(-\d{4})?$", "99999 or 99999-9999"),
    "United Kingdom": (r"^[A-Z]{1,2}[0-9][0-9A-Z]?\s?[0-9][A-Z]{2}$", "Format: AA99 9AA or A9 9AA"),
    "Japan": (r"^\d{3}-?\d{4}$", "7 digits, dash after 3rd optional"),
    "Australia": (r"^\d{4}$", "4 digits"),
    "India": (r"^\d{6}$", "6 digits"),
    "Germany": (r"^\d{5}$", "5 digits"),
    "Canada": (r"^[A-Z]\d[A-Z]\s?\d[A-Z]\d$", "Format: A1A 1A1"),
    "France": (r"^\d{5}$", "5 digits"),
    "Ireland": (r"^(D\d{1,2}|D6W|[A-Z]\d{2}\s?[A-Z0-9]{4})$", "Eircode or D+digits"),
    "Singapore": (r"^\d{6}$", "6 digits"),
    "Netherlands": (r"^\d{4}\s?[A-Z]{2}$", "4 digits + 2 letters"),
    "Brazil": (r"^\d{5}(-?\d{3})?$", "5 or 8 digits"),
    "Spain": (r"^\d{5}$", "5 digits"),
    "Switzerland": (r"^\d{4}$", "4 digits"),
    "Italy": (r"^\d{5}$", "5 digits"),
    "Sweden": (r"^\d{3}\s?\d{2}$", "5 digits, space after 3rd optional"),
    "Israel": (r"^\d{5}(\d{2})?$", "5 or 7 digits"),
    "Belgium": (r"^\d{4}$", "4 digits"),
    "Mexico": (r"^\d{5}$", "5 digits"),
    "New Zealand": (r"^\d{4}$", "4 digits"),
    "China": (r"^\d{6}$", "6 digits"),
    "Denmark": (r"^\d{4}$", "4 digits"),
    "Argentina": (r"^[A-Z]\d{4}([A-Z]{3})?$", "1 letter + 4 digits + optional 3 letters"),
    "Portugal": (r"^\d{4}-?\d{3}$", "7 digits, dash after 4th optional"),
    "Austria": (r"^\d{4}$", "4 digits"),
    "Finland": (r"^\d{5}$", "5 digits"),
    "Taiwan": (r"^\d{3,6}$", "3-6 digits"),
    "Norway": (r"^\d{4}$", "4 digits"),
    "South Africa": (r"^([BS]-?)?\d{4}$", "4 digits, optional B- or S- prefix"),
    "Korea, Republic of": (r"^\d{5,6}$", "5 or 6 digits"),
    "Thailand": (r"^\d{5}$", "5 digits"),
    "Colombia": (r"^\d{6}$", "6 digits"),
    "Morocco": (r"^\d{5}$", "5 digits"),
    "Luxembourg": (r"^\d{4}$", "4 digits"),
    "Indonesia": (r"^\d{5}$", "5 digits"),
    "Chile": (r"^\d{7}$|^\d{3}-\d{4}$", "7 digits or NNN-NNNN"),
    "Malaysia": (r"^\d{5}$", "5 digits"),
}


# ---------------------------------------------------------------------------
# BANKING RULES
# ---------------------------------------------------------------------------

BANKING_RULES = {
    "United States": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "N/A"
    },
    "United Kingdom": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "Optional", "iban": "Required", "name_on_account": "Optional"
    },
    "India": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Optional"
    },
    "Japan": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Required",
        "branch_id": "Required", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Australia": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Canada": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Required", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "N/A"
    },
    "Germany": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "Optional"
    },
    "Singapore": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Required", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Brazil": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Required", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "Required"
    },
    "France": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Optional", "account_number": "Optional", "bic": "Optional",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "Optional"
    },
    "Italy": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Optional", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "N/A"
    },
    "Israel": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Required", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "N/A"
    },
    "New Zealand": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Optional"
    },
    "Switzerland": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "Optional"
    },
    "Spain": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Optional", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "Optional"
    },
    "Ireland": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "Required"
    },
    "Netherlands": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "Optional"
    },
    "Argentina": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "China": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Required", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Optional"
    },
    "Korea, Republic of": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Mexico": {
        "routing_transit_number": "N/A", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Optional"
    },
    "Sweden": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "Optional"
    },
    "Belgium": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "N/A"
    },
    "Hong Kong": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Required", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Denmark": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Required", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "N/A"
    },
    "Thailand": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Optional", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "South Africa": {
        "routing_transit_number": "N/A", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Required", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Austria": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "Optional"
    },
    "Taiwan": {
        "routing_transit_number": "Required", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Optional", "account_number": "Required", "bic": "Optional",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Required"
    },
    "Norway": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Optional", "bic": "Optional",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "Optional"
    },
    "Poland": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Optional", "bic": "Optional",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "Optional"
    },
    "Colombia": {
        "routing_transit_number": "N/A", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "N/A", "account_number": "Required", "bic": "N/A",
        "roll_number": "N/A", "iban": "N/A", "name_on_account": "Optional"
    },
    "United Arab Emirates": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "Optional",
        "branch_id": "Optional", "account_number": "Optional", "bic": "Optional",
        "roll_number": "N/A", "iban": "Optional", "name_on_account": "Optional"
    },
    "Portugal": {
        "routing_transit_number": "Optional", "bank_name": "Required", "branch_name": "N/A",
        "branch_id": "Optional", "account_number": "Optional", "bic": "Required",
        "roll_number": "N/A", "iban": "Required", "name_on_account": "N/A"
    },
}



# ---------------------------------------------------------------------------
# COUNTRY ALIASES
# ---------------------------------------------------------------------------

COUNTRY_ALIASES = {
    "US": "United States of America",
    "USA": "United States of America",
    "UNITED STATES": "United States of America",
    "UNITED STATES OF AMERICA": "United States of America",
    "UK": "United Kingdom",
    "GB": "United Kingdom",
    "GBR": "United Kingdom",
    "GREAT BRITAIN": "United Kingdom",
    "ENGLAND": "United Kingdom",
    "UNITED KINGDOM": "United Kingdom",
    "AU": "Australia",
    "AUS": "Australia",
    "AUSTRALIA": "Australia",
    "CA": "Canada",
    "CAN": "Canada",
    "CANADA": "Canada",
    "DE": "Germany",
    "DEU": "Germany",
    "GERMANY": "Germany",
    "FR": "France",
    "FRA": "France",
    "FRANCE": "France",
    "IN": "India",
    "IND": "India",
    "INDIA": "India",
    "JP": "Japan",
    "JPN": "Japan",
    "JAPAN": "Japan",
    "SG": "Singapore",
    "SGP": "Singapore",
    "SINGAPORE": "Singapore",
    "BR": "Brazil",
    "BRA": "Brazil",
    "BRAZIL": "Brazil",
    "NL": "Netherlands",
    "NLD": "Netherlands",
    "NETHERLANDS": "Netherlands",
    "THE NETHERLANDS": "Netherlands",
    "IT": "Italy",
    "ITA": "Italy",
    "ITALY": "Italy",
    "ES": "Spain",
    "ESP": "Spain",
    "SPAIN": "Spain",
    "CH": "Switzerland",
    "CHE": "Switzerland",
    "SWITZERLAND": "Switzerland",
    "SE": "Sweden",
    "SWE": "Sweden",
    "SWEDEN": "Sweden",
    "BE": "Belgium",
    "BEL": "Belgium",
    "BELGIUM": "Belgium",
    "MX": "Mexico",
    "MEX": "Mexico",
    "MEXICO": "Mexico",
    "AR": "Argentina",
    "ARG": "Argentina",
    "ARGENTINA": "Argentina",
    "CN": "China",
    "CHN": "China",
    "CHINA": "China",
    "KR": "Korea, Republic of",
    "KOR": "Korea, Republic of",
    "KOREA": "Korea, Republic of",
    "SOUTH KOREA": "Korea, Republic of",
    "KOREA, REPUBLIC OF": "Korea, Republic of",
    "REPUBLIC OF KOREA": "Korea, Republic of",
    "HK": "Hong Kong",
    "HKG": "Hong Kong",
    "HONG KONG": "Hong Kong",
    "DK": "Denmark",
    "DNK": "Denmark",
    "DENMARK": "Denmark",
    "NO": "Norway",
    "NOR": "Norway",
    "NORWAY": "Norway",
    "FI": "Finland",
    "FIN": "Finland",
    "FINLAND": "Finland",
    "AT": "Austria",
    "AUT": "Austria",
    "AUSTRIA": "Austria",
    "PT": "Portugal",
    "PRT": "Portugal",
    "PORTUGAL": "Portugal",
    "IE": "Ireland",
    "IRL": "Ireland",
    "IRELAND": "Ireland",
    "LU": "Luxembourg",
    "LUX": "Luxembourg",
    "LUXEMBOURG": "Luxembourg",
    "TW": "Taiwan",
    "TWN": "Taiwan",
    "TAIWAN": "Taiwan",
    "TH": "Thailand",
    "THA": "Thailand",
    "THAILAND": "Thailand",
    "NZ": "New Zealand",
    "NZL": "New Zealand",
    "NEW ZEALAND": "New Zealand",
    "IL": "Israel",
    "ISR": "Israel",
    "ISRAEL": "Israel",
    "ZA": "South Africa",
    "ZAF": "South Africa",
    "SOUTH AFRICA": "South Africa",
    "CO": "Colombia",
    "COL": "Colombia",
    "COLOMBIA": "Colombia",
    "MA": "Morocco",
    "MAR": "Morocco",
    "MOROCCO": "Morocco",
    "KY": "Cayman Islands",
    "CYM": "Cayman Islands",
    "CAYMAN ISLANDS": "Cayman Islands",
    "AE": "United Arab Emirates",
    "ARE": "United Arab Emirates",
    "UAE": "United Arab Emirates",
    "UNITED ARAB EMIRATES": "United Arab Emirates",
}

# Secondary alias table for BANKING_RULES keys (different canonical form than ADDRESS_RULES)
_BANKING_COUNTRY_ALIASES = {
    "UNITED STATES OF AMERICA": "United States",
    "US": "United States",
    "USA": "United States",
    "UNITED STATES": "United States",
    "UK": "United Kingdom",
    "GB": "United Kingdom",
    "GBR": "United Kingdom",
    "GREAT BRITAIN": "United Kingdom",
    "ENGLAND": "United Kingdom",
    "UNITED KINGDOM": "United Kingdom",
    "AU": "Australia",
    "AUS": "Australia",
    "AUSTRALIA": "Australia",
    "CA": "Canada",
    "CAN": "Canada",
    "CANADA": "Canada",
    "DE": "Germany",
    "DEU": "Germany",
    "GERMANY": "Germany",
    "FR": "France",
    "FRA": "France",
    "FRANCE": "France",
    "IN": "India",
    "IND": "India",
    "INDIA": "India",
    "JP": "Japan",
    "JPN": "Japan",
    "JAPAN": "Japan",
    "SG": "Singapore",
    "SGP": "Singapore",
    "SINGAPORE": "Singapore",
    "BR": "Brazil",
    "BRA": "Brazil",
    "BRAZIL": "Brazil",
    "NL": "Netherlands",
    "NLD": "Netherlands",
    "NETHERLANDS": "Netherlands",
    "THE NETHERLANDS": "Netherlands",
    "IT": "Italy",
    "ITA": "Italy",
    "ITALY": "Italy",
    "ES": "Spain",
    "ESP": "Spain",
    "SPAIN": "Spain",
    "CH": "Switzerland",
    "CHE": "Switzerland",
    "SWITZERLAND": "Switzerland",
    "SE": "Sweden",
    "SWE": "Sweden",
    "SWEDEN": "Sweden",
    "BE": "Belgium",
    "BEL": "Belgium",
    "BELGIUM": "Belgium",
    "MX": "Mexico",
    "MEX": "Mexico",
    "MEXICO": "Mexico",
    "AR": "Argentina",
    "ARG": "Argentina",
    "ARGENTINA": "Argentina",
    "CN": "China",
    "CHN": "China",
    "CHINA": "China",
    "KR": "Korea, Republic of",
    "KOR": "Korea, Republic of",
    "KOREA": "Korea, Republic of",
    "SOUTH KOREA": "Korea, Republic of",
    "REPUBLIC OF KOREA": "Korea, Republic of",
    "KOREA, REPUBLIC OF": "Korea, Republic of",
    "HK": "Hong Kong",
    "HKG": "Hong Kong",
    "HONG KONG": "Hong Kong",
    "DK": "Denmark",
    "DNK": "Denmark",
    "DENMARK": "Denmark",
    "NO": "Norway",
    "NOR": "Norway",
    "NORWAY": "Norway",
    "FI": "Finland",
    "FIN": "Finland",
    "FINLAND": "Finland",
    "AT": "Austria",
    "AUT": "Austria",
    "AUSTRIA": "Austria",
    "PT": "Portugal",
    "PRT": "Portugal",
    "PORTUGAL": "Portugal",
    "IE": "Ireland",
    "IRL": "Ireland",
    "IRELAND": "Ireland",
    "LU": "Luxembourg",
    "LUX": "Luxembourg",
    "LUXEMBOURG": "Luxembourg",
    "TW": "Taiwan",
    "TWN": "Taiwan",
    "TAIWAN": "Taiwan",
    "TH": "Thailand",
    "THA": "Thailand",
    "THAILAND": "Thailand",
    "NZ": "New Zealand",
    "NZL": "New Zealand",
    "NEW ZEALAND": "New Zealand",
    "IL": "Israel",
    "ISR": "Israel",
    "ISRAEL": "Israel",
    "ZA": "South Africa",
    "ZAF": "South Africa",
    "SOUTH AFRICA": "South Africa",
    "CO": "Colombia",
    "COL": "Colombia",
    "COLOMBIA": "Colombia",
    "MA": "Morocco",
    "MAR": "Morocco",
    "MOROCCO": "Morocco",
    "KY": "Cayman Islands",
    "CYM": "Cayman Islands",
    "CAYMAN ISLANDS": "Cayman Islands",
    "AE": "United Arab Emirates",
    "ARE": "United Arab Emirates",
    "UAE": "United Arab Emirates",
    "UNITED ARAB EMIRATES": "United Arab Emirates",
    "PL": "Poland",
    "POL": "Poland",
    "POLAND": "Poland",
}

# ISO-2 lookup for Workday output column
_COUNTRY_ISO2 = {
    "United States of America": "US",
    "United Kingdom": "GB",
    "Australia": "AU",
    "Canada": "CA",
    "Germany": "DE",
    "France": "FR",
    "India": "IN",
    "Japan": "JP",
    "Singapore": "SG",
    "Brazil": "BR",
    "Netherlands": "NL",
    "Italy": "IT",
    "Spain": "ES",
    "Switzerland": "CH",
    "Sweden": "SE",
    "Belgium": "BE",
    "Mexico": "MX",
    "Argentina": "AR",
    "China": "CN",
    "Korea, Republic of": "KR",
    "Hong Kong": "HK",
    "Denmark": "DK",
    "Norway": "NO",
    "Finland": "FI",
    "Austria": "AT",
    "Portugal": "PT",
    "Ireland": "IE",
    "Luxembourg": "LU",
    "Taiwan": "TW",
    "Thailand": "TH",
    "New Zealand": "NZ",
    "Israel": "IL",
    "South Africa": "ZA",
    "Colombia": "CO",
    "Morocco": "MA",
    "Cayman Islands": "KY",
    "United Arab Emirates": "AE",
    "Poland": "PL",
    "Indonesia": "ID",
    "Chile": "CL",
    "Malaysia": "MY",
}

# Countries where postal code is uppercased for Workday load
_UPPERCASE_POSTAL_COUNTRIES = {"United States of America", "United Kingdom", "Canada"}



# ---------------------------------------------------------------------------
# NORMALIZE COUNTRY
# ---------------------------------------------------------------------------

def normalize_country(raw: str) -> str:
    """
    Strip whitespace, uppercase, look up in COUNTRY_ALIASES.
    Return the canonical name if found, else return the original stripped value.
    """
    if not isinstance(raw, str):
        return str(raw).strip() if raw is not None else ""
    stripped = raw.strip()
    upper = stripped.upper()
    return COUNTRY_ALIASES.get(upper, stripped)


def _normalize_banking_country(raw: str) -> str:
    """
    Resolve a country string to the BANKING_RULES canonical key.
    Checks _BANKING_COUNTRY_ALIASES first, then falls back to direct BANKING_RULES key lookup.
    """
    if not isinstance(raw, str):
        return str(raw).strip() if raw is not None else ""
    stripped = raw.strip()
    upper = stripped.upper()
    candidate = _BANKING_COUNTRY_ALIASES.get(upper, stripped)
    if candidate in BANKING_RULES:
        return candidate
    addr_canonical = COUNTRY_ALIASES.get(upper, stripped)
    if addr_canonical.upper() in _BANKING_COUNTRY_ALIASES:
        return _BANKING_COUNTRY_ALIASES[addr_canonical.upper()]
    if addr_canonical in BANKING_RULES:
        return addr_canonical
    return stripped


# ---------------------------------------------------------------------------
# COLUMN DETECTION - ADDRESS
# ---------------------------------------------------------------------------

_ADDRESS_COL_KEYWORDS = {
    "address_line_1": ["address line 1", "line 1", "addr1", "address1", "street address", "street"],
    "address_line_2": ["address line 2", "line 2", "addr2", "address2"],
    "address_line_3": ["address line 3", "line 3", "addr3"],
    "city": ["city", "municipality", "town"],
    "city_subdivision_1": ["city subdivision", "neighborhood", "neighbourhood", "district", "suburb"],
    "region": ["region", "state", "province", "county", "prefecture", "land"],
    "region_subdivision_1": ["region subdivision", "post office", "district"],
    "postal_code": ["postal code", "zip code", "zip", "postcode", "post code"],
    "country": ["country name", "country code", "country"],
}


def detect_address_columns(df_columns: list) -> dict:
    """
    Map dataframe columns to standard internal address keys.
    Returns {internal_key: actual_column_name}. Only includes matched keys.
    """
    result = {}
    lower_cols = {col: col.lower().strip() for col in df_columns}
    used_cols = set()

    for internal_key, keywords in _ADDRESS_COL_KEYWORDS.items():
        best_match = None
        best_score = -1
        for col, col_lower in lower_cols.items():
            if col in used_cols:
                continue
            for kw in keywords:
                if col_lower == kw:
                    score = len(kw) + 100
                elif kw in col_lower:
                    score = len(kw)
                else:
                    continue
                if score > best_score:
                    best_score = score
                    best_match = col
        if best_match is not None:
            result[internal_key] = best_match
            used_cols.add(best_match)

    return result


# ---------------------------------------------------------------------------
# COLUMN DETECTION - BANKING
# ---------------------------------------------------------------------------

_BANKING_COL_KEYWORDS = {
    "routing_transit_number": [
        "routing transit number", "routing number", "transit number", "aba number", "aba",
        "sort code", "sort", "routing", "transit"
    ],
    "bank_name": ["bank name", "bank"],
    "branch_name": ["branch name", "branch"],
    "branch_id": ["branch id", "branch code", "branch number"],
    "account_number": ["account number", "account no", "account #", "account"],
    "bic": ["bic code", "swift code", "bic", "swift"],
    "roll_number": ["roll number", "roll no", "roll #", "roll"],
    "iban": ["iban"],
    "name_on_account": ["name on account", "account name", "account holder"],
    "country": ["country name", "country code", "country"],
}


def detect_banking_columns(df_columns: list) -> dict:
    """
    Map dataframe columns to standard internal banking keys.
    Returns {internal_key: actual_column_name}. Only includes matched keys.
    """
    result = {}
    lower_cols = {col: col.lower().strip() for col in df_columns}
    used_cols = set()

    for internal_key, keywords in _BANKING_COL_KEYWORDS.items():
        best_match = None
        best_score = -1
        for col, col_lower in lower_cols.items():
            if col in used_cols:
                continue
            for kw in keywords:
                if col_lower == kw:
                    score = len(kw) + 100
                elif kw in col_lower:
                    score = len(kw)
                else:
                    continue
                if score > best_score:
                    best_score = score
                    best_match = col
        if best_match is not None:
            result[internal_key] = best_match
            used_cols.add(best_match)

    return result


# ---------------------------------------------------------------------------
# HELPER UTILITIES
# ---------------------------------------------------------------------------

def _is_blank(value) -> bool:
    """Return True if value is None, NaN, or empty/whitespace string."""
    if value is None:
        return True
    if isinstance(value, float):
        import math
        return math.isnan(value)
    return str(value).strip() == ""


def _str_val(value) -> str:
    """Return stripped string value, or empty string if blank."""
    if _is_blank(value):
        return ""
    return str(value).strip()


def _read_dataframe(file_path: str):
    """Read CSV or Excel file into a pandas DataFrame."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext in (".xlsx", ".xls", ".xlsm", ".xlsb"):
        return pd.read_excel(file_path, dtype=str)
    try:
        return pd.read_csv(file_path, dtype=str, encoding="utf-8")
    except UnicodeDecodeError:
        return pd.read_csv(file_path, dtype=str, encoding="latin-1")


def _validate_postal_code(postal_code: str, country: str):
    """
    Validate postal code against pattern for country.
    Returns (True, None) if valid or no pattern exists.
    Returns (False, expected_format_hint) if invalid.
    """
    if country not in POSTAL_CODE_PATTERNS:
        return True, None
    pattern, hint = POSTAL_CODE_PATTERNS[country]
    normalized = postal_code.strip().upper()
    if re.match(pattern, normalized, re.IGNORECASE):
        return True, None
    return False, hint


def _validate_iban(iban: str):
    """
    Basic IBAN format check: 2 uppercase letters + 2 digits + alphanumeric, length 15-34.
    Returns (True, None) or (False, reason_string).
    """
    s = iban.strip().upper().replace(" ", "")
    pattern = r"^[A-Z]{2}\d{2}[A-Z0-9]{11,30}$"
    if not (15 <= len(s) <= 34):
        return False, f"IBAN length must be 15-34 characters (got {len(s)})"
    if not re.match(pattern, s):
        return False, "IBAN must start with 2 letters + 2 digits followed by alphanumeric characters"
    return True, None


def _validate_bic(bic: str):
    """
    Basic BIC/SWIFT check: 8 or 11 alphanumeric characters.
    Returns (True, None) or (False, reason_string).
    """
    s = bic.strip().upper().replace(" ", "")
    if len(s) not in (8, 11):
        return False, f"BIC must be 8 or 11 characters (got {len(s)})"
    if not re.match(r"^[A-Z0-9]+$", s):
        return False, "BIC must contain only letters and digits"
    return True, None



# ---------------------------------------------------------------------------
# VALIDATE ADDRESS FILE
# ---------------------------------------------------------------------------

def validate_address_file(file_path: str, output_path: Optional[str] = None) -> dict:
    """
    Read a CSV or Excel file, validate each row against Workday address rules,
    and write an annotated Excel output.

    Returns a results dict with keys:
        total, valid_count, warning_count, error_count,
        rows (list of per-row dicts), output_path, column_map, errors (file-level),
        country_summary
    """
    results = {
        "total": 0,
        "valid_count": 0,
        "warning_count": 0,
        "error_count": 0,
        "rows": [],
        "output_path": None,
        "column_map": {},
        "errors": [],
        "country_summary": {},
    }

    if pd is None:
        results["errors"].append(
            "pandas is not installed. Install with: pip install pandas openpyxl"
        )
        return results

    try:
        df = _read_dataframe(file_path)
    except Exception as exc:
        results["errors"].append(f"Failed to read file '{file_path}': {exc}")
        return results

    col_map = detect_address_columns(list(df.columns))
    results["column_map"] = col_map
    results["total"] = len(df)

    for idx, row in df.iterrows():
        row_num = idx + 2  # 1-based, accounting for header row
        issues = []
        suggestions = []

        raw_country = (
            _str_val(row.get(col_map.get("country", ""), ""))
            if "country" in col_map
            else ""
        )
        canonical_country = normalize_country(raw_country) if raw_country else ""
        rules = ADDRESS_RULES.get(canonical_country, _ADDRESS_RULES_GENERIC)
        unknown_country = canonical_country not in ADDRESS_RULES and canonical_country != ""

        if unknown_country:
            suggestions.append(
                f"Country '{raw_country}' not found in Workday address rules; generic rules applied."
            )

        corrected = {}

        def get_field(key):
            col = col_map.get(key)
            if col is None:
                return ""
            return _str_val(row.get(col, ""))

        for field, requirement in rules.items():
            value = get_field(field)
            corrected[field] = value

            if requirement == "Required":
                if _is_blank(value):
                    issues.append(f"ERROR: '{field}' is required but missing.")
            elif requirement == "Not Allowed":
                if not _is_blank(value):
                    issues.append(
                        f"WARNING: '{field}' is not allowed for "
                        f"{canonical_country or 'this country'} but has value '{value}'."
                    )
                    suggestions.append(f"Remove value from '{field}' field.")

        # Postal code format validation
        if "postal_code" in col_map and not _is_blank(get_field("postal_code")):
            pc = get_field("postal_code")
            if canonical_country in POSTAL_CODE_PATTERNS:
                valid_pc, hint = _validate_postal_code(pc, canonical_country)
                if not valid_pc:
                    issues.append(
                        f"WARNING: Postal code '{pc}' does not match expected format for "
                        f"{canonical_country}. Expected: {hint}."
                    )
                    suggestions.append(f"Correct postal code format to: {hint}")

        # Workday-ready values
        wd_addr1 = corrected.get("address_line_1", "")
        wd_city = corrected.get("city", "")
        wd_region = corrected.get("region", "")
        wd_postal = corrected.get("postal_code", "")
        wd_iso2 = _COUNTRY_ISO2.get(canonical_country, "")

        if canonical_country in _UPPERCASE_POSTAL_COUNTRIES and wd_postal:
            wd_postal = wd_postal.upper()

        error_issues = [i for i in issues if i.startswith("ERROR")]
        warning_issues = [i for i in issues if i.startswith("WARNING")]

        if error_issues:
            status = "error"
            results["error_count"] += 1
        elif warning_issues:
            status = "warning"
            results["warning_count"] += 1
        else:
            status = "valid"
            results["valid_count"] += 1

        summary_key = canonical_country or "(unknown)"
        if summary_key not in results["country_summary"]:
            results["country_summary"][summary_key] = {
                "total": 0, "valid": 0, "warning": 0, "error": 0
            }
        results["country_summary"][summary_key]["total"] += 1
        results["country_summary"][summary_key][status] += 1

        row_dict = {
            "row_num": row_num,
            "original": {col: _str_val(row.get(col, "")) for col in df.columns},
            "status": status,
            "issues": issues,
            "suggestions": suggestions,
            "corrected": corrected,
            "workday": {
                "address_line_1": wd_addr1,
                "city": wd_city,
                "region": wd_region,
                "postal_code": wd_postal,
                "country_iso2": wd_iso2,
            },
            "raw_country": raw_country,
            "canonical_country": canonical_country,
        }
        results["rows"].append(row_dict)

    if output_path is None:
        fd, output_path = tempfile.mkstemp(
            suffix=".xlsx", prefix="workday_address_validation_"
        )
        os.close(fd)

    try:
        generate_address_excel(results, output_path, df)
        results["output_path"] = output_path
    except Exception as exc:
        results["errors"].append(f"Failed to write output Excel: {exc}")

    return results


# ---------------------------------------------------------------------------
# VALIDATE BANKING FILE
# ---------------------------------------------------------------------------

def validate_banking_file(file_path: str, output_path: Optional[str] = None) -> dict:
    """
    Read a CSV or Excel file, validate each row against Workday banking rules,
    and write an annotated Excel output.

    Returns a results dict with keys:
        total, valid_count, warning_count, error_count,
        rows (list of per-row dicts), output_path, column_map, errors (file-level),
        country_summary
    """
    results = {
        "total": 0,
        "valid_count": 0,
        "warning_count": 0,
        "error_count": 0,
        "rows": [],
        "output_path": None,
        "column_map": {},
        "errors": [],
        "country_summary": {},
    }

    if pd is None:
        results["errors"].append(
            "pandas is not installed. Install with: pip install pandas openpyxl"
        )
        return results

    try:
        df = _read_dataframe(file_path)
    except Exception as exc:
        results["errors"].append(f"Failed to read file '{file_path}': {exc}")
        return results

    col_map = detect_banking_columns(list(df.columns))
    results["column_map"] = col_map
    results["total"] = len(df)

    for idx, row in df.iterrows():
        row_num = idx + 2
        issues = []
        suggestions = []

        def get_field(key):
            col = col_map.get(key)
            if col is None:
                return ""
            return _str_val(row.get(col, ""))

        raw_country = get_field("country") if "country" in col_map else ""
        banking_country = _normalize_banking_country(raw_country) if raw_country else ""
        rules = BANKING_RULES.get(banking_country)
        unknown_country = False

        if rules is None:
            unknown_country = True
            rules = {}
            if raw_country:
                suggestions.append(
                    f"Country '{raw_country}' not found in Workday banking rules; "
                    "no field-level validation applied."
                )

        all_banking_fields = [
            "routing_transit_number", "bank_name", "branch_name", "branch_id",
            "account_number", "bic", "roll_number", "iban", "name_on_account",
        ]

        for field in all_banking_fields:
            requirement = rules.get(field, "Optional")
            value = get_field(field)

            if requirement == "Required":
                if _is_blank(value):
                    issues.append(
                        f"ERROR: '{field}' is required for "
                        f"{banking_country or 'this country'} but missing."
                    )
            elif requirement == "N/A":
                if not _is_blank(value):
                    issues.append(
                        f"WARNING: '{field}' is not applicable for "
                        f"{banking_country or 'this country'} but has value '{value}'."
                    )
                    suggestions.append(
                        f"Remove value from '{field}' — not used for this country."
                    )

            # Format checks
            if not _is_blank(value):
                if field == "iban":
                    valid_iban, iban_reason = _validate_iban(value)
                    if not valid_iban:
                        issues.append(f"WARNING: IBAN format invalid — {iban_reason}.")
                        suggestions.append(
                            "Correct IBAN to standard format: CC99XXXXXX... (15-34 chars)."
                        )
                elif field == "bic":
                    valid_bic, bic_reason = _validate_bic(value)
                    if not valid_bic:
                        issues.append(f"WARNING: BIC/SWIFT format invalid — {bic_reason}.")
                        suggestions.append(
                            "Correct BIC to 8 or 11 alphanumeric characters."
                        )

        error_issues = [i for i in issues if i.startswith("ERROR")]
        warning_issues = [i for i in issues if i.startswith("WARNING")]

        if error_issues:
            status = "error"
            results["error_count"] += 1
        elif warning_issues:
            status = "warning"
            results["warning_count"] += 1
        else:
            status = "valid"
            results["valid_count"] += 1

        workday_ready = "YES" if status in ("valid", "warning") and not error_issues else "NO"

        summary_key = banking_country or raw_country or "(unknown)"
        if summary_key not in results["country_summary"]:
            results["country_summary"][summary_key] = {
                "total": 0, "valid": 0, "warning": 0, "error": 0
            }
        results["country_summary"][summary_key]["total"] += 1
        results["country_summary"][summary_key][status] += 1

        row_dict = {
            "row_num": row_num,
            "original": {col: _str_val(row.get(col, "")) for col in df.columns},
            "status": status,
            "issues": issues,
            "suggestions": suggestions,
            "raw_country": raw_country,
            "banking_country": banking_country,
            "workday_ready": workday_ready,
        }
        results["rows"].append(row_dict)

    if output_path is None:
        fd, output_path = tempfile.mkstemp(
            suffix=".xlsx", prefix="workday_banking_validation_"
        )
        os.close(fd)

    try:
        generate_banking_excel(results, output_path, df)
        results["output_path"] = output_path
    except Exception as exc:
        results["errors"].append(f"Failed to write output Excel: {exc}")

    return results



# ---------------------------------------------------------------------------
# EXCEL OUTPUT - ADDRESS
# ---------------------------------------------------------------------------

def generate_address_excel(results: dict, output_path: str, source_df=None):
    """
    Write address validation results to an Excel workbook.

    Sheet 1 Validation Results: original columns + status/issues/suggestions + Workday columns.
    Sheet 2 Summary: totals and country breakdown.
    """
    if openpyxl is None:
        raise ImportError(
            "openpyxl is not installed. Install with: pip install openpyxl"
        )

    wb = openpyxl.Workbook()

    fill_header = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    fill_valid = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    fill_warning = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    fill_error = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    fill_summary_header = PatternFill(
        start_color="2C5F8A", end_color="2C5F8A", fill_type="solid"
    )

    font_header = Font(bold=True, color="FFFFFF", size=11)
    font_summary_header = Font(bold=True, color="FFFFFF", size=11)

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws = wb.active
    ws.title = "Validation Results"

    if source_df is not None:
        orig_cols = list(source_df.columns)
    elif results["rows"]:
        orig_cols = list(results["rows"][0]["original"].keys())
    else:
        orig_cols = []

    extra_cols = [
        "Validation Status", "Issues", "Suggestions",
        "Workday Address Line 1", "Workday City", "Workday Region",
        "Workday Postal Code", "Workday Country ISO2",
    ]
    all_headers = orig_cols + extra_cols

    for col_idx, header in enumerate(all_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 30

    _status_fill = {"valid": fill_valid, "warning": fill_warning, "error": fill_error}
    _status_label = {"valid": "VALID", "warning": "WARNING", "error": "ERROR"}

    for row_data in results["rows"]:
        excel_row = row_data["row_num"]
        row_fill = _status_fill.get(row_data["status"], fill_valid)

        for col_idx, col_name in enumerate(orig_cols, start=1):
            val = row_data["original"].get(col_name, "")
            cell = ws.cell(row=excel_row, column=col_idx, value=val)
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=False)

        base = len(orig_cols)

        cell = ws.cell(
            row=excel_row, column=base + 1,
            value=_status_label.get(row_data["status"], "")
        )
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="top")

        issues_str = "; ".join(row_data["issues"]) if row_data["issues"] else ""
        cell = ws.cell(row=excel_row, column=base + 2, value=issues_str)
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

        suggestions_str = "; ".join(row_data["suggestions"]) if row_data["suggestions"] else ""
        cell = ws.cell(row=excel_row, column=base + 3, value=suggestions_str)
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

        wd = row_data.get("workday", {})
        workday_values = [
            wd.get("address_line_1", ""),
            wd.get("city", ""),
            wd.get("region", ""),
            wd.get("postal_code", ""),
            wd.get("country_iso2", ""),
        ]
        for i, val in enumerate(workday_values, start=4):
            cell = ws.cell(row=excel_row, column=base + i, value=val)
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top")

    for col_idx, header in enumerate(all_headers, start=1):
        max_len = len(str(header))
        for row_data in results["rows"]:
            cell = ws.cell(row=row_data["row_num"], column=col_idx)
            cell_val = str(cell.value) if cell.value else ""
            max_len = max(max_len, min(len(cell_val), 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, max_len + 2)

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet(title="Summary")

    for col_idx, header in enumerate(["Metric", "Value"], start=1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = fill_summary_header
        cell.font = font_summary_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    summary_rows = [
        ("Total Rows", results["total"]),
        ("Valid", results["valid_count"]),
        ("Warnings", results["warning_count"]),
        ("Errors", results["error_count"]),
    ]
    for r_idx, (label, value) in enumerate(summary_rows, start=2):
        ws2.cell(row=r_idx, column=1, value=label).border = thin_border
        ws2.cell(row=r_idx, column=2, value=value).border = thin_border

    if results.get("country_summary"):
        country_hdr_row = len(summary_rows) + 3
        for col_idx, header in enumerate(
            ["Country", "Total", "Valid", "Warnings", "Errors"], start=1
        ):
            cell = ws2.cell(row=country_hdr_row, column=col_idx, value=header)
            cell.fill = fill_summary_header
            cell.font = font_summary_header
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        r_idx = country_hdr_row + 1
        for country, counts in sorted(results["country_summary"].items()):
            ws2.cell(row=r_idx, column=1, value=country).border = thin_border
            ws2.cell(row=r_idx, column=2, value=counts["total"]).border = thin_border
            ws2.cell(row=r_idx, column=3, value=counts["valid"]).border = thin_border
            ws2.cell(row=r_idx, column=4, value=counts["warning"]).border = thin_border
            ws2.cell(row=r_idx, column=5, value=counts["error"]).border = thin_border
            r_idx += 1

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 15
    ws2.column_dimensions["D"].width = 15
    ws2.column_dimensions["E"].width = 15

    try:
        wb.save(output_path)
    except Exception as exc:
        raise IOError(f"Failed to save Excel file to '{output_path}': {exc}") from exc



# ---------------------------------------------------------------------------
# EXCEL OUTPUT - BANKING
# ---------------------------------------------------------------------------

def generate_banking_excel(results: dict, output_path: str, source_df=None):
    """
    Write banking validation results to an Excel workbook.

    Sheet 1 Validation Results: original columns + status/issues/suggestions + Workday Ready.
    Sheet 2 Summary: totals and country breakdown.
    """
    if openpyxl is None:
        raise ImportError(
            "openpyxl is not installed. Install with: pip install openpyxl"
        )

    wb = openpyxl.Workbook()

    fill_header = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    fill_valid = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    fill_warning = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    fill_error = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    fill_summary_header = PatternFill(
        start_color="2C5F8A", end_color="2C5F8A", fill_type="solid"
    )

    font_header = Font(bold=True, color="FFFFFF", size=11)
    font_summary_header = Font(bold=True, color="FFFFFF", size=11)

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws = wb.active
    ws.title = "Validation Results"

    if source_df is not None:
        orig_cols = list(source_df.columns)
    elif results["rows"]:
        orig_cols = list(results["rows"][0]["original"].keys())
    else:
        orig_cols = []

    extra_cols = ["Validation Status", "Issues", "Suggestions", "Workday Ready"]
    all_headers = orig_cols + extra_cols

    for col_idx, header in enumerate(all_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 30

    _status_fill = {"valid": fill_valid, "warning": fill_warning, "error": fill_error}
    _status_label = {"valid": "VALID", "warning": "WARNING", "error": "ERROR"}

    for row_data in results["rows"]:
        excel_row = row_data["row_num"]
        row_fill = _status_fill.get(row_data["status"], fill_valid)

        for col_idx, col_name in enumerate(orig_cols, start=1):
            val = row_data["original"].get(col_name, "")
            cell = ws.cell(row=excel_row, column=col_idx, value=val)
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=False)

        base = len(orig_cols)

        cell = ws.cell(
            row=excel_row, column=base + 1,
            value=_status_label.get(row_data["status"], "")
        )
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="top")

        issues_str = "; ".join(row_data["issues"]) if row_data["issues"] else ""
        cell = ws.cell(row=excel_row, column=base + 2, value=issues_str)
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

        suggestions_str = "; ".join(row_data["suggestions"]) if row_data["suggestions"] else ""
        cell = ws.cell(row=excel_row, column=base + 3, value=suggestions_str)
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

        cell = ws.cell(
            row=excel_row, column=base + 4,
            value=row_data.get("workday_ready", "NO")
        )
        cell.fill = row_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="top")

    for col_idx, header in enumerate(all_headers, start=1):
        max_len = len(str(header))
        for row_data in results["rows"]:
            cell = ws.cell(row=row_data["row_num"], column=col_idx)
            cell_val = str(cell.value) if cell.value else ""
            max_len = max(max_len, min(len(cell_val), 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, max_len + 2)

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet(title="Summary")

    for col_idx, header in enumerate(["Metric", "Value"], start=1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = fill_summary_header
        cell.font = font_summary_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    summary_rows = [
        ("Total Rows", results["total"]),
        ("Valid", results["valid_count"]),
        ("Warnings", results["warning_count"]),
        ("Errors", results["error_count"]),
    ]
    for r_idx, (label, value) in enumerate(summary_rows, start=2):
        ws2.cell(row=r_idx, column=1, value=label).border = thin_border
        ws2.cell(row=r_idx, column=2, value=value).border = thin_border

    if results.get("country_summary"):
        country_hdr_row = len(summary_rows) + 3
        for col_idx, header in enumerate(
            ["Country", "Total", "Valid", "Warnings", "Errors"], start=1
        ):
            cell = ws2.cell(row=country_hdr_row, column=col_idx, value=header)
            cell.fill = fill_summary_header
            cell.font = font_summary_header
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        r_idx = country_hdr_row + 1
        for country, counts in sorted(results["country_summary"].items()):
            ws2.cell(row=r_idx, column=1, value=country).border = thin_border
            ws2.cell(row=r_idx, column=2, value=counts["total"]).border = thin_border
            ws2.cell(row=r_idx, column=3, value=counts["valid"]).border = thin_border
            ws2.cell(row=r_idx, column=4, value=counts["warning"]).border = thin_border
            ws2.cell(row=r_idx, column=5, value=counts["error"]).border = thin_border
            r_idx += 1

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 15
    ws2.column_dimensions["D"].width = 15
    ws2.column_dimensions["E"].width = 15

    try:
        wb.save(output_path)
    except Exception as exc:
        raise IOError(f"Failed to save Excel file to '{output_path}': {exc}") from exc


# ---------------------------------------------------------------------------
# QUICK SMOKE TEST (run directly: python workday_validator.py)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    print("workday_validator.py loaded successfully.")
    print(f"  ADDRESS_RULES countries : {len(ADDRESS_RULES)}")
    print(f"  POSTAL_CODE_PATTERNS    : {len(POSTAL_CODE_PATTERNS)}")
    print(f"  BANKING_RULES countries : {len(BANKING_RULES)}")
    print(f"  COUNTRY_ALIASES entries : {len(COUNTRY_ALIASES)}")

    tests = [
        ("US", "United States of America"),
        ("usa", "United States of America"),
        ("UK", "United Kingdom"),
        ("GBR", "United Kingdom"),
        ("DE", "Germany"),
        ("KR", "Korea, Republic of"),
        ("AE", "United Arab Emirates"),
        ("Unknown Country XYZ", "Unknown Country XYZ"),
    ]
    print("\nnormalize_country tests:")
    all_pass = True
    for raw, expected in tests:
        result = normalize_country(raw)
        status = "PASS" if result == expected else "FAIL"
        if status == "FAIL":
            all_pass = False
        print(f"  [{status}] normalize_country({raw!r}) -> {result!r}  (expected {expected!r})")

    pc_tests = [
        ("United States of America", "90210", True),
        ("United States of America", "90210-1234", True),
        ("United States of America", "9021", False),
        ("United Kingdom", "SW1A 1AA", True),
        ("Germany", "10115", True),
        ("Germany", "1015", False),
        ("Canada", "M5V 3L9", True),
        ("Japan", "100-0001", True),
        ("Japan", "10001", False),
    ]
    print("\nPostal code validation tests:")
    for country, pc, expected_valid in pc_tests:
        valid, hint = _validate_postal_code(pc, country)
        status = "PASS" if valid == expected_valid else "FAIL"
        if status == "FAIL":
            all_pass = False
        print(
            f"  [{status}] {country} / {pc!r} -> valid={valid}  (expected {expected_valid})"
        )

    iban_tests = [
        ("GB29NWBK60161331926819", True),
        ("DE89370400440532013000", True),
        ("INVALID", False),
        ("GB29", False),
    ]
    print("\nIBAN validation tests:")
    for iban, expected_valid in iban_tests:
        valid, reason = _validate_iban(iban)
        status = "PASS" if valid == expected_valid else "FAIL"
        if status == "FAIL":
            all_pass = False
        print(
            f"  [{status}] {iban!r} -> valid={valid}"
            + (f"  ({reason})" if reason else "")
        )

    bic_tests = [
        ("NWBKGB2L", True),
        ("NWBKGB2LXXX", True),
        ("NWBK", False),
        ("NWBKGB2L!!", False),
    ]
    print("\nBIC validation tests:")
    for bic, expected_valid in bic_tests:
        valid, reason = _validate_bic(bic)
        status = "PASS" if valid == expected_valid else "FAIL"
        if status == "FAIL":
            all_pass = False
        print(
            f"  [{status}] {bic!r} -> valid={valid}"
            + (f"  ({reason})" if reason else "")
        )

    print(f"\nAll tests passed: {all_pass}")
